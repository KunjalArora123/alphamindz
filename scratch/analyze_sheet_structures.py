import os
import openpyxl
import glob
import sys

sys.stdout.reconfigure(encoding='utf-8')

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = sorted(glob.glob(os.path.join(alpha_dir, "*.xlsx")))

for fpath in excel_files:
    fname = os.path.basename(fpath)
    print("\n" + "="*80)
    print(f"FILE: {fname}")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- SHEET: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---")
        
        # Collect header row (rows 1-6)
        print("Header / First 6 rows:")
        for r in range(1, min(7, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
            if any(row_vals):
                print(f"  Row {r:02d}: {row_vals}")
                
        # Sample question row (row 7 or 8)
        print("Sample Data rows:")
        for r in range(7, min(12, ws.max_row + 1)):
            row_vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
            if any(row_vals):
                print(f"  Row {r:02d}: {row_vals}")
                
        # Image anchors info
        images = ws._images if hasattr(ws, '_images') else []
        if images:
            print(f"Images count: {len(images)}")
            sample_imgs = []
            for img in images[:5]:
                r = None
                c = None
                if hasattr(img.anchor, '_from'):
                    r = img.anchor._from.row + 1
                    c = img.anchor._from.col + 1
                elif hasattr(img.anchor, 'row'):
                    r = img.anchor.row
                    c = img.anchor.col
                sample_imgs.append(f"Row {r}, Col {c}")
            print(f"  First 5 image anchors: {sample_imgs}")

