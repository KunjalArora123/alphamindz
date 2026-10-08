import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r"d:\xampp\htdocs\AlphaMindz\alpha_tests\Qusetion Paper 8th AAT 2026 (3).xlsx"
wb = openpyxl.load_workbook(fpath, data_only=True)

ws = wb['8th Numerical '] # note space at end
print(f"Sheet: '{ws.title}' max_row={ws.max_row}")

images = ws._images if hasattr(ws, '_images') else []
print(f"Total images in openpyxl: {len(images)}")
for idx, img in enumerate(images):
    r, c = None, None
    if hasattr(img.anchor, '_from'):
        r = img.anchor._from.row + 1
        c = img.anchor._from.col + 1
    elif hasattr(img.anchor, 'row'):
        r = img.anchor.row
        c = img.anchor.col
    print(f"  Img #{idx}: row={r}, col={c}")

print("\n--- ROWS & QUESTIONS ---")
for r in range(1, ws.max_row + 1):
    val1 = ws.cell(r, 1).value
    val2 = ws.cell(r, 2).value
    val3 = ws.cell(r, 3).value
    val4 = ws.cell(r, 4).value
    val5 = ws.cell(r, 5).value
    val6 = ws.cell(r, 6).value
    val7 = ws.cell(r, 7).value
    val8 = ws.cell(r, 8).value
    val9 = ws.cell(r, 9).value
    
    non_empty = [c for c in [val1, val2, val3, val4, val5, val6, val7, val8, val9] if c is not None and str(c).strip() != '']
    if non_empty:
        print(f"Row {r:02d}: col1={val1}, col2={val2}, col3='{str(val3)[:40] if val3 else ''}'")

