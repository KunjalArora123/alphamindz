import os
import openpyxl
import glob
from PIL import Image
import io

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = sorted(glob.glob(os.path.join(alpha_dir, "*.xlsx")))

for fpath in excel_files:
    fname = os.path.basename(fpath)
    print("==========================================================")
    print("EXCEL FILE:", fname)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- SHEET: {sheet_name} ---")
        
        # Check images in openpyxl worksheet
        images = ws._images if hasattr(ws, '_images') else []
        print(f"Total images in sheet: {len(images)}")
        
        # Map images to rows if possible
        img_by_row = {}
        for idx, img in enumerate(images):
            # openpyxl image anchor can be OneCellAnchor or TwoCellAnchor
            row = None
            if hasattr(img.anchor, '_from'):
                row = img.anchor._from.row + 1 # 1-indexed row
                col = img.anchor._from.col + 1
            elif hasattr(img.anchor, 'row'):
                row = img.anchor.row
                col = img.anchor.col
            
            if row not in img_by_row:
                img_by_row[row] = []
            img_by_row[row].append((idx, img))
            
        print(f"Image rows found: {sorted(list(img_by_row.keys()))}")
        
        # Look through rows
        question_rows = []
        for r in range(1, ws.max_row + 1):
            val1 = ws.cell(r, 1).value
            val2 = ws.cell(r, 2).value
            val3 = ws.cell(r, 3).value
            val4 = ws.cell(r, 4).value
            val5 = ws.cell(r, 5).value
            val6 = ws.cell(r, 6).value
            val7 = ws.cell(r, 7).value
            val8 = ws.cell(r, 8).value
            val9 = ws.cell(r, 9).value
            
            # Check if this row looks like a header or a question row
            r_str = f"R{r:02d}: col1={val1}, col2={val2}, col3={str(val3)[:40] if val3 else ''}"
            # Check if images present near row r
            imgs_here = img_by_row.get(r, [])
            if imgs_here:
                r_str += f" | IMAGES AT ROW {r}: {len(imgs_here)}"
            
            # Print non-empty rows
            non_empty = [c for c in [val1, val2, val3, val4, val5, val6, val7, val8, val9] if c is not None and str(c).strip() != '']
            if non_empty or imgs_here:
                print(r_str)

