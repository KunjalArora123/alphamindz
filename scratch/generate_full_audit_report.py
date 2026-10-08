import os
import openpyxl
import glob
import sys
import json
import re

sys.stdout.reconfigure(encoding='utf-8')

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = sorted(glob.glob(os.path.join(alpha_dir, "*.xlsx")))

report = []

for fpath in excel_files:
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        
        # Build image map for this sheet: (row, col) -> image_obj list
        img_map = {}
        images = ws._images if hasattr(ws, '_images') else []
        for idx, img in enumerate(images):
            r = None
            c = None
            if hasattr(img.anchor, '_from'):
                r = img.anchor._from.row + 1
                c = img.anchor._from.col + 1
            elif hasattr(img.anchor, 'row'):
                r = img.anchor.row
                c = img.anchor.col
            
            if r is not None and c is not None:
                if (r, c) not in img_map:
                    img_map[(r, c)] = []
                img_map[(r, c)].append(img)
                
        # Find all question rows
        for r in range(1, ws.max_row + 1):
            val1 = ws.cell(r, 1).value
            val2 = ws.cell(r, 2).value
            
            q_num = None
            for candidate in [val2, val1]:
                if candidate is not None:
                    try:
                        val_str = str(candidate).strip()
                        if val_str.isdigit() and 1 <= int(val_str) <= 30:
                            q_num = int(val_str)
                            break
                    except ValueError:
                        pass
            
            if q_num is not None:
                val3 = ws.cell(r, 3).value # Q text / image
                val4 = ws.cell(r, 4).value # Opt A
                val5 = ws.cell(r, 5).value # Opt B
                val6 = ws.cell(r, 6).value # Opt C
                val7 = ws.cell(r, 7).value # Opt D
                val8 = ws.cell(r, 8).value # Fig/Graph or Answer
                val9 = ws.cell(r, 9).value # Answer
                val10 = ws.cell(r, 10).value # Answer / extra
                
                # Determine correct option
                ans = None
                for candidate_ans in [val9, val8, val10]:
                    if candidate_ans is not None:
                        c_str = str(candidate_ans).strip().upper()
                        if c_str in ['A', 'B', 'C', 'D', 'E'] or re.match(r'^[A-E]+$', c_str):
                            ans = c_str
                            break
                
                # Collect images near row r (rows r-1, r, r+1)
                row_imgs = {}
                for check_r in [r-1, r, r+1]:
                    for c in range(1, 12):
                        if (check_r, c) in img_map:
                            row_imgs[c] = row_imgs.get(c, 0) + len(img_map[(check_r, c)])
                            
                report.append({
                    'file': fname,
                    'sheet': sname,
                    'row': r,
                    'q_num': q_num,
                    'q_text': str(val3).strip() if val3 is not None else '',
                    'opt_a': str(val4).strip() if val4 is not None else '',
                    'opt_b': str(val5).strip() if val5 is not None else '',
                    'opt_c': str(val6).strip() if val6 is not None else '',
                    'opt_d': str(val7).strip() if val7 is not None else '',
                    'val8': str(val8).strip() if val8 is not None else '',
                    'val9': str(val9).strip() if val9 is not None else '',
                    'ans': ans,
                    'img_cols': row_imgs
                })

with open(r"d:\xampp\htdocs\AlphaMindz\scratch\detailed_question_report.json", "w", encoding="utf-8") as f:
    json.dump(report, f, indent=2, ensure_ascii=False)

print(f"Generated report for {len(report)} questions.")

