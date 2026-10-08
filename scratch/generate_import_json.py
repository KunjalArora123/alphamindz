import os
import openpyxl
import glob
import sys
import json
from PIL import Image
import io
import re

sys.stdout.reconfigure(encoding='utf-8')

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
output_img_dir = r"d:\xampp\htdocs\AlphaMindz\questions-images"
os.makedirs(output_img_dir, exist_ok=True)

excel_files_info = [
    {
        'file': 'Qusetion Paper 10th AAT 2026.xlsx',
        'title': '10th AAT 2026 Assessment',
        'slug': '10th-aat-2026-assessment',
        'key': '10th'
    },
    {
        'file': 'Qusetion Paper 8th AAT 2026 (3).xlsx',
        'title': '8th AAT 2026 (3) Assessment',
        'slug': '8th-aat-2026-3-assessment',
        'key': '8th'
    },
    {
        'file': 'Qusetion Paper 9th AAT 2026.xlsx',
        'title': '9th AAT 2026 Assessment',
        'slug': '9th-aat-2026-assessment',
        'key': '9th'
    },
    {
        'file': 'Qusetion Paper ARTS 2026.xlsx',
        'title': 'ARTS 2026 Assessment',
        'slug': 'arts-2026-assessment',
        'key': 'arts'
    },
    {
        'file': 'Qusetion Paper COMMERCE 2026.xlsx',
        'title': 'COMMERCE 2026 Assessment',
        'slug': 'commerce-2026-assessment',
        'key': 'commerce'
    },
    {
        'file': 'Qusetion Paper SCIENCE 2026.xlsx',
        'title': 'SCIENCE 2026 Assessment',
        'slug': 'science-2026-assessment',
        'key': 'science'
    }
]

import_data = []

for info in excel_files_info:
    fpath = os.path.join(alpha_dir, info['file'])
    if not os.path.exists(fpath):
        print(f"ERROR: File {fpath} not found!")
        continue
        
    wb = openpyxl.load_workbook(fpath, data_only=True)
    assessment_item = {
        'title': info['title'],
        'slug': info['slug'],
        'description': f"Comprehensive aptitude assessment for {info['title']}.",
        'time_limit': 45,
        'parts': []
    }
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        clean_part_name = sname.strip()
        sheet_slug = clean_part_name.lower().replace(' ', '_').replace('\t', '')
        
        part_item = {
            'part_name': clean_part_name,
            'description': f"{clean_part_name} section evaluating core competencies.",
            'questions': []
        }
        
        # Save images in worksheet
        sheet_imgs = ws._images if hasattr(ws, '_images') else []
        img_map = {} # (row, col) -> list of relative filepaths
        
        for idx, img in enumerate(sheet_imgs):
            r, c = None, None
            if hasattr(img.anchor, '_from'):
                r = img.anchor._from.row + 1
                c = img.anchor._from.col + 1
            elif hasattr(img.anchor, 'row'):
                r = img.anchor.row
                c = img.anchor.col
                
            if r is None or c is None:
                continue
                
            img_bytes = img._data()
            im = Image.open(io.BytesIO(img_bytes))
            ext = im.format.lower() if im.format else 'png'
            if ext == 'jpeg': ext = 'jpg'
            
            img_filename = f"{info['key']}_{sheet_slug}_r{r}_c{c}_{idx}.{ext}"
            img_rel_path = f"questions-images/{img_filename}"
            full_out_path = os.path.join(output_img_dir, img_filename)
            
            with open(full_out_path, 'wb') as f_out:
                f_out.write(img_bytes)
                
            if (r, c) not in img_map:
                img_map[(r, c)] = []
            img_map[(r, c)].append(img_rel_path)
            
        # Parse question rows
        question_rows = []
        for r in range(1, ws.max_row + 1):
            val1 = ws.cell(r, 1).value
            val2 = ws.cell(r, 2).value
            
            q_num = None
            for candidate in [val2, val1]:
                if candidate is not None:
                    try:
                        v_str = str(candidate).strip()
                        if v_str.isdigit() and 1 <= int(v_str) <= 30:
                            q_num = int(v_str)
                            break
                    except ValueError:
                        pass
            
            if q_num is not None:
                question_rows.append((r, q_num))
                
        for r, q_num in question_rows:
            val3 = ws.cell(r, 3).value # Q text / image
            val4 = ws.cell(r, 4).value # Opt A
            val5 = ws.cell(r, 5).value # Opt B
            val6 = ws.cell(r, 6).value # Opt C
            val7 = ws.cell(r, 7).value # Opt D
            val8 = ws.cell(r, 8).value # Fig/Graph or Answer
            val9 = ws.cell(r, 9).value # Answer
            val10 = ws.cell(r, 10).value
            
            q_text = str(val3).strip() if val3 is not None else ''
            opt_a = str(val4).strip() if val4 is not None else ''
            opt_b = str(val5).strip() if val5 is not None else ''
            opt_c = str(val6).strip() if val6 is not None else ''
            opt_d = str(val7).strip() if val7 is not None else ''
            
            # Determine answer
            ans = ''
            for candidate_ans in [val9, val8, val10]:
                if candidate_ans is not None:
                    c_str = str(candidate_ans).strip().upper()
                    if c_str in ['A', 'B', 'C', 'D', 'E'] or re.match(r'^[A-Z1-9]+$', c_str):
                        ans = c_str
                        break
                        
            # Collect images for this question row (rows r-1, r, r+1)
            image_path = ''
            image_path_2 = ''
            opt_a_image = ''
            opt_b_image = ''
            opt_c_image = ''
            opt_d_image = ''
            
            for check_r in [r-1, r, r+1]:
                for c in [3, 8, 9]:
                    if (check_r, c) in img_map:
                        for p in img_map[(check_r, c)]:
                            if not image_path:
                                image_path = p
                            elif not image_path_2 and p != image_path:
                                image_path_2 = p
                                
                if (check_r, 4) in img_map and not opt_a_image:
                    opt_a_image = img_map[(check_r, 4)][0]
                if (check_r, 5) in img_map and not opt_b_image:
                    opt_b_image = img_map[(check_r, 5)][0]
                if (check_r, 6) in img_map and not opt_c_image:
                    opt_c_image = img_map[(check_r, 6)][0]
                if (check_r, 7) in img_map and not opt_d_image:
                    opt_d_image = img_map[(check_r, 7)][0]
                    
            part_item['questions'].append({
                'question_number': q_num,
                'question_text': q_text,
                'option_a': opt_a,
                'option_a_image': opt_a_image,
                'option_b': opt_b,
                'option_b_image': opt_b_image,
                'option_c': opt_c,
                'option_c_image': opt_c_image,
                'option_d': opt_d,
                'option_d_image': opt_d_image,
                'image_path': image_path,
                'image_path_2': image_path_2,
                'correct_option': ans
            })
            
        assessment_item['parts'].append(part_item)
        
    import_data.append(assessment_item)

out_json_path = r"d:\xampp\htdocs\AlphaMindz\scratch\full_test_import.json"
with open(out_json_path, 'w', encoding='utf-8') as f:
    json.dump(import_data, f, indent=2, ensure_ascii=False)

print(f"Successfully generated {out_json_path} for {len(import_data)} assessments!")

