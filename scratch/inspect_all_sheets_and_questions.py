import os
import openpyxl
import glob
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = sorted(glob.glob(os.path.join(alpha_dir, "*.xlsx")))

all_data = {}

for fpath in excel_files:
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    all_data[fname] = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_info = {
            'max_row': ws.max_row,
            'max_col': ws.max_column,
            'images': [],
            'questions': []
        }
        
        # Check images and their locations
        images = ws._images if hasattr(ws, '_images') else []
        for idx, img in enumerate(images):
            row = None
            col = None
            if hasattr(img.anchor, '_from'):
                row = img.anchor._from.row + 1
                col = img.anchor._from.col + 1
            elif hasattr(img.anchor, 'row'):
                row = img.anchor.row
                col = img.anchor.col
            
            sheet_info['images'].append({
                'index': idx,
                'row': row,
                'col': col,
                'width': img.width if hasattr(img, 'width') else None,
                'height': img.height if hasattr(img, 'height') else None
            })
            
        # Parse questions (rows where col 2 is integer question_number or row contains question content)
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # remove trailing Nones
            while row_vals and row_vals[-1] is None:
                row_vals.pop()
            if not row_vals:
                continue
                
            q_num = ws.cell(r, 2).value
            # Check if q_num is integer or digit string
            is_q = False
            try:
                if q_num is not None:
                    int(str(q_num).strip())
                    is_q = True
            except ValueError:
                pass
                
            sheet_info['questions'].append({
                'row': r,
                'is_q_row': is_q,
                'row_values': [str(v) if v is not None else '' for v in row_vals]
            })
            
        all_data[fname][sheet_name] = sheet_info

with open(r"d:\xampp\htdocs\AlphaMindz\scratch\excel_inspection.json", "w", encoding="utf-8") as f:
    json.dump(all_data, f, indent=2, ensure_ascii=False)

print("Inspection completed. Summary:")
for fname, sheets in all_data.items():
    print(f"\n{fname}:")
    for sname, sinfo in sheets.items():
        q_count = sum(1 for q in sinfo['questions'] if q['is_q_row'])
        img_count = len(sinfo['images'])
        print(f"  [{sname}]: {q_count} questions detected, {img_count} images in openpyxl")

