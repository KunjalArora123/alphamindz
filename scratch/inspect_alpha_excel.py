import os
import openpyxl
import zipfile
import glob

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = glob.glob(os.path.join(alpha_dir, "*.xlsx"))

print(f"Found {len(excel_files)} Excel files:")
for f in excel_files:
    print("\n=========================================")
    print("FILE:", os.path.basename(f))
    wb = openpyxl.load_workbook(f, data_only=True)
    print("Sheets:", wb.sheetnames)
    
    # Check images in zip
    with zipfile.ZipFile(f, 'r') as z:
        media_files = [item for item in z.namelist() if item.startswith('xl/media/')]
        print(f"Embedded images in zip ({len(media_files)}):", media_files[:10])
    
    for sname in wb.sheetnames:
        sheet = wb[sname]
        print(f"\n--- Sheet: {sname} (max_row={sheet.max_row}, max_column={sheet.max_column}) ---")
        img_count = len(sheet._images) if hasattr(sheet, '_images') else 0
        print(f"Sheet images via openpyxl: {img_count}")
        
        # Print first 5 rows
        for r in range(1, min(6, sheet.max_row + 1)):
            row_vals = [sheet.cell(r, c).value for c in range(1, min(12, sheet.max_column + 1))]
            if any(row_vals):
                print(f"  Row {r}: {row_vals}")

