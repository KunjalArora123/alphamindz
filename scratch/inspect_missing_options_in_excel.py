import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

fpath = r"d:\xampp\htdocs\AlphaMindz\alpha_tests\Qusetion Paper ARTS 2026.xlsx"
wb = openpyxl.load_workbook(fpath, data_only=True)

ws = wb['ARTS VERBAL ABILITY']
print("=== ARTS VERBAL ABILITY Q18, Q19 ===")
for r in [24, 25]: # rows 24, 25 in Excel
    vals = [ws.cell(r, c).value for c in range(1, 10)]
    print(f"Row {r}: {vals}")

ws2 = wb['SPATIAL ABILITY']
print("\n=== SPATIAL ABILITY Q11..Q20 ===")
for r in range(15, 25):
    vals = [ws2.cell(r, c).value for c in range(1, 10)]
    print(f"Row {r:02d}: {vals}")

