import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

arts_file = r"d:\xampp\htdocs\AlphaMindz\alpha_tests\Qusetion Paper ARTS 2026.xlsx"
comm_file = r"d:\xampp\htdocs\AlphaMindz\alpha_tests\Qusetion Paper COMMERCE 2026.xlsx"

print("================ ARTS VISUAL ==================")
wb = openpyxl.load_workbook(arts_file, data_only=True)
ws = wb['ARTS VISUAL']
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
images = ws._images if hasattr(ws, '_images') else []
print(f"Images count: {len(images)}")
for img in images[:15]:
    r = img.anchor._from.row + 1 if hasattr(img.anchor, '_from') else img.anchor.row
    c = img.anchor._from.col + 1 if hasattr(img.anchor, '_from') else img.anchor.col
    print(f"  Img anchor: row {r}, col {c}")

for r in range(1, min(25, ws.max_row + 1)):
    row_vals = [ws.cell(r, c).value for c in range(1, min(12, ws.max_column + 1))]
    print(f"Row {r:02d}: {row_vals}")

print("\n================ COMMERCE CLERICAL ==================")
wb2 = openpyxl.load_workbook(comm_file, data_only=True)
ws2 = wb2['CLERICAL ABILITY']
print(f"Max row: {ws2.max_row}, Max col: {ws2.max_column}")
images2 = ws2._images if hasattr(ws2, '_images') else []
print(f"Images count: {len(images2)}")
for img in images2[:15]:
    r = img.anchor._from.row + 1 if hasattr(img.anchor, '_from') else img.anchor.row
    c = img.anchor._from.col + 1 if hasattr(img.anchor, '_from') else img.anchor.col
    print(f"  Img anchor: row {r}, col {c}")

for r in range(1, min(25, ws2.max_row + 1)):
    row_vals = [ws2.cell(r, c).value for c in range(1, min(12, ws2.max_column + 1))]
    print(f"Row {r:02d}: {row_vals}")

