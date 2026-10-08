import os
import openpyxl
import glob
import sys
import json
import mysql.connector
from PIL import Image
import io
import re

sys.stdout.reconfigure(encoding='utf-8')

# Database connection details
DB_CONFIG = {
    'host': 'localhost',
    'user': 'root',
    'password': '',
    'database': 'alphamindz',
    'charset': 'utf8mb4'
}

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
output_img_dir = r"d:\xampp\htdocs\AlphaMindz\questions-images"
os.makedirs(output_img_dir, exist_ok=True)

excel_files_info = [
    {
        'file': 'Qusetion Paper 10th AAT 2026.xlsx',
        'title': '10th AAT 2026 Assessment',
        'slug': '10th-aat-2026-assessment',
        'key': '10th'
    },
    {
        'file': 'Qusetion Paper 8th AAT 2026 (3).xlsx',
        'title': '8th AAT 2026 (3) Assessment',
        'slug': '8th-aat-2026-3-assessment',
        'key': '8th'
    },
    {
        'file': 'Qusetion Paper 9th AAT 2026.xlsx',
        'title': '9th AAT 2026 Assessment',
        'slug': '9th-aat-2026-assessment',
        'key': '9th'
    },
    {
        'file': 'Qusetion Paper ARTS 2026.xlsx',
        'title': 'ARTS 2026 Assessment',
        'slug': 'arts-2026-assessment',
        'key': 'arts'
    },
    {
        'file': 'Qusetion Paper COMMERCE 2026.xlsx',
        'title': 'COMMERCE 2026 Assessment',
        'slug': 'commerce-2026-assessment',
        'key': 'commerce'
    },
    {
        'file': 'Qusetion Paper SCIENCE 2026.xlsx',
        'title': 'SCIENCE 2026 Assessment',
        'slug': 'science-2026-assessment',
        'key': 'science'
    }
]

def run_seeder():
    conn = mysql.connector.connect(**DB_CONFIG)
    cursor = conn.cursor(dictionary=True)
    
    print("=== STEP 1: ALTER TABLE SCHEMA IF NEEDED ===")
    cursor.execute("ALTER TABLE questions MODIFY COLUMN correct_option VARCHAR(20) NOT NULL")
    conn.commit()
    print("Altered `questions.correct_option` to VARCHAR(20)")

    print("\n=== STEP 2: DELETE EXISTING APTITUDE TESTS ===")
    # Delete assessments other than MBTI (33) & Interest (34)
    cursor.execute("SELECT id, title FROM assessments WHERE id NOT IN (33, 34)")
    old_assessments = cursor.fetchall()
    old_ids = [a['id'] for a in old_assessments]
    
    if old_ids:
        old_ids_str = ', '.join(str(i) for i in old_ids)
        print(f"Deleting old assessments IDs: {old_ids_str}")
        cursor.execute(f"DELETE FROM test_answers WHERE question_id IN (SELECT id FROM questions WHERE assessment_id IN ({old_ids_str}))")
        cursor.execute(f"DELETE FROM questions WHERE assessment_id IN ({old_ids_str})")
        cursor.execute(f"DELETE FROM assessment_parts WHERE assessment_id IN ({old_ids_str})")
        cursor.execute(f"DELETE FROM assessments WHERE id IN ({old_ids_str})")
        conn.commit()
        print("Deleted old assessment data successfully!")
    else:
        print("No old aptitude assessments found to delete.")

    print("\n=== STEP 3: RE-UPLOAD ALL 6 ASSESSMENTS ===")
    
    total_assessments_created = 0
    total_parts_created = 0
    total_questions_created = 0
    total_images_extracted = 0

    for info in excel_files_info:
        fpath = os.path.join(alpha_dir, info['file'])
        if not os.path.exists(fpath):
            print(f"ERROR: File {fpath} does not exist!")
            continue
            
        print(f"\n---------------------------------------------------------")
        print(f"Processing File: {info['file']} ({info['title']})")
        
        # Insert assessment
        cursor.execute(
            "INSERT INTO assessments (title, slug, description, time_limit, status, created_at, updated_at) "
            "VALUES (%s, %s, %s, %s, 'active', NOW(), NOW())",
            (info['title'], info['slug'], f"Comprehensive assessment for {info['title']}.", 45)
        )
        assessment_id = cursor.lastrowid
        total_assessments_created += 1
        print(f"Created Assessment ID: {assessment_id}")
        
        wb = openpyxl.load_workbook(fpath, data_only=True)
        
        for sname in wb.sheetnames:
            ws = wb[sname]
            clean_part_name = sname.strip()
            sheet_slug = clean_part_name.lower().replace(' ', '_').replace('\t', '')
            
            cursor.execute(
                "INSERT INTO assessment_parts (assessment_id, part_name, description, created_at, updated_at) "
                "VALUES (%s, %s, %s, NOW(), NOW())",
                (assessment_id, clean_part_name, f"{clean_part_name} section evaluating core competencies.")
            )
            part_id = cursor.lastrowid
            total_parts_created += 1
            
            # Map images in sheet: (row, col) -> list of saved img paths
            sheet_imgs = ws._images if hasattr(ws, '_images') else []
            img_map = {} # (row, col) -> list of relative filepaths
            
            for idx, img in enumerate(sheet_imgs):
                r, c = None, None
                if hasattr(img.anchor, '_from'):
                    r = img.anchor._from.row + 1
                    c = img.anchor._from.col + 1
                elif hasattr(img.anchor, 'row'):
                    r = img.anchor.row
                    c = img.anchor.col
                    
                if r is None or c is None:
                    continue
                    
                img_bytes = img._data()
                im = Image.open(io.BytesIO(img_bytes))
                ext = im.format.lower() if im.format else 'png'
                if ext == 'jpeg': ext = 'jpg'
                
                img_filename = f"{info['key']}_{sheet_slug}_r{r}_c{c}_{idx}.{ext}"
                img_rel_path = f"questions-images/{img_filename}"
                full_out_path = os.path.join(output_img_dir, img_filename)
                
                with open(full_out_path, 'wb') as f_out:
                    f_out.write(img_bytes)
                total_images_extracted += 1
                
                if (r, c) not in img_map:
                    img_map[(r, c)] = []
                img_map[(r, c)].append(img_rel_path)
                
            # Process question rows in sheet
            question_rows = []
            for r in range(1, ws.max_row + 1):
                val1 = ws.cell(r, 1).value
                val2 = ws.cell(r, 2).value
                
                q_num = None
                for candidate in [val2, val1]:
                    if candidate is not None:
                        try:
                            v_str = str(candidate).strip()
                            if v_str.isdigit() and 1 <= int(v_str) <= 30:
                                q_num = int(v_str)
                                break
                        except ValueError:
                            pass
                
                if q_num is not None:
                    question_rows.append((r, q_num))
                    
            print(f"  Part '{clean_part_name}' (ID {part_id}): Found {len(question_rows)} questions, {len(sheet_imgs)} images.")
            
            for r, q_num in question_rows:
                val3 = ws.cell(r, 3).value # Q text / image
                val4 = ws.cell(r, 4).value # Opt A
                val5 = ws.cell(r, 5).value # Opt B
                val6 = ws.cell(r, 6).value # Opt C
                val7 = ws.cell(r, 7).value # Opt D
                val8 = ws.cell(r, 8).value # Fig/Graph or Answer
                val9 = ws.cell(r, 9).value # Answer
                val10 = ws.cell(r, 10).value
                
                q_text = str(val3).strip() if val3 is not None else ''
                opt_a = str(val4).strip() if val4 is not None else ''
                opt_b = str(val5).strip() if val5 is not None else ''
                opt_c = str(val6).strip() if val6 is not None else ''
                opt_d = str(val7).strip() if val7 is not None else ''
                
                # Determine answer
                ans = ''
                for candidate_ans in [val9, val8, val10]:
                    if candidate_ans is not None:
                        c_str = str(candidate_ans).strip().upper()
                        if c_str in ['A', 'B', 'C', 'D', 'E'] or re.match(r'^[A-Z1-9]+$', c_str):
                            ans = c_str
                            break
                            
                # Collect images for this question row (checking rows r-1, r, r+1)
                image_path = ''
                image_path_2 = ''
                opt_a_image = ''
                opt_b_image = ''
                opt_c_image = ''
                opt_d_image = ''
                
                for check_r in [r-1, r, r+1]:
                    # Col 3 (C) or Col 8 (H) or Col 9 (I) -> main question image
                    for c in [3, 8, 9]:
                        if (check_r, c) in img_map:
                            for p in img_map[(check_r, c)]:
                                if not image_path:
                                    image_path = p
                                elif not image_path_2 and p != image_path:
                                    image_path_2 = p
                                    
                    # Col 4 -> Opt A
                    if (check_r, 4) in img_map and not opt_a_image:
                        opt_a_image = img_map[(check_r, 4)][0]
                    # Col 5 -> Opt B
                    if (check_r, 5) in img_map and not opt_b_image:
                        opt_b_image = img_map[(check_r, 5)][0]
                    # Col 6 -> Opt C
                    if (check_r, 6) in img_map and not opt_c_image:
                        opt_c_image = img_map[(check_r, 6)][0]
                    # Col 7 -> Opt D
                    if (check_r, 7) in img_map and not opt_d_image:
                        opt_d_image = img_map[(check_r, 7)][0]
                        
                # Insert into questions table
                cursor.execute(
                    "INSERT INTO questions ("
                    "  assessment_id, part_id, subject, question_number, question_text, "
                    "  option_a, option_a_image, option_b, option_b_image, option_c, option_c_image, "
                    "  option_d, option_d_image, image_path, image_path_2, correct_option, "
                    "  created_at, updated_at"
                    ") VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW(), NOW())",
                    (
                        assessment_id, part_id, clean_part_name, q_num, q_text,
                        opt_a, opt_a_image, opt_b, opt_b_image, opt_c, opt_c_image,
                        opt_d, opt_d_image, image_path, image_path_2, ans
                    )
                )
                total_questions_created += 1

        conn.commit()

    conn.close()
    
    print("\n==========================================================")
    print("SUMMARY OF SEEDING COMPLETED:")
    print(f"  Assessments Created: {total_assessments_created}")
    print(f"  Assessment Parts Created: {total_parts_created}")
    print(f"  Questions Created: {total_questions_created}")
    print(f"  Images Extracted: {total_images_extracted}")
    print("==========================================================")

if __name__ == '__main__':
    run_seeder()

