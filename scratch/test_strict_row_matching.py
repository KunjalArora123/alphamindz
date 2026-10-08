import openpyxl
import os
import glob
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
excel_files = sorted(glob.glob(os.path.join(alpha_dir, "*.xlsx")))

for fpath in excel_files:
    fname = os.path.basename(fpath)
    print(f"\n==========================================")
    print(f"FILE: {fname}")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        
        # Build image list
        images = ws._images if hasattr(ws, '_images') else []
        img_by_row = {}
        for idx, img in enumerate(images):
            r, c = None, None
            if hasattr(img.anchor, '_from'):
                r = img.anchor._from.row + 1
                c = img.anchor._from.col + 1
            elif hasattr(img.anchor, 'row'):
                r = img.anchor.row
                c = img.anchor.col
            if r is not None and c is not None:
                if r not in img_by_row:
                    img_by_row[r] = []
                img_by_row[r].append((c, idx, img))
                
        # Question rows map: row -> q_num
        q_rows = {}
        for r in range(1, ws.max_row + 1):
            val1 = ws.cell(r, 1).value
            val2 = ws.cell(r, 2).value
            for candidate in [val2, val1]:
                if candidate is not None:
                    try:
                        v_str = str(candidate).strip()
                        if v_str.isdigit() and 1 <= int(v_str) <= 30:
                            q_rows[r] = int(v_str)
                            break
                    except ValueError:
                        pass
                        
        # Check images that don't match any q_row exactly
        unmatched_imgs = []
        matched_imgs = []
        for r, img_list in img_by_row.items():
            if r in q_rows:
                matched_imgs.extend(img_list)
            else:
                unmatched_imgs.append((r, img_list))
                
        print(f"  Sheet [{sname}]: Total Images={len(images)}, Strictly Matched to Question Row={len(matched_imgs)}, Unmatched={len(unmatched_imgs)}")
        if unmatched_imgs:
            for r, ulist in unmatched_imgs:
                print(f"     UNMATCHED Image at Row {r}: {len(ulist)} images. Nearest Q rows: {[qr for qr in q_rows.keys() if abs(qr-r)<=2]}")

