import os
import openpyxl
import glob
import sys
import json
import mysql.connector

sys.stdout.reconfigure(encoding='utf-8')

DB_CONFIG = {
    'host': 'localhost',
    'user': 'root',
    'password': '',
    'database': 'alphamindz',
    'charset': 'utf8mb4'
}

alpha_dir = r"d:\xampp\htdocs\AlphaMindz\alpha_tests"
fcpath = r"d:\xampp\htdocs\AlphaMindz"

excel_files_info = [
    {
        'file': 'Qusetion Paper 10th AAT 2026.xlsx',
        'title': '10th AAT 2026 Assessment',
        'slug': '10th-aat-2026-assessment'
    },
    {
        'file': 'Qusetion Paper 8th AAT 2026 (3).xlsx',
        'title': '8th AAT 2026 (3) Assessment',
        'slug': '8th-aat-2026-3-assessment'
    },
    {
        'file': 'Qusetion Paper 9th AAT 2026.xlsx',
        'title': '9th AAT 2026 Assessment',
        'slug': '9th-aat-2026-assessment'
    },
    {
        'file': 'Qusetion Paper ARTS 2026.xlsx',
        'title': 'ARTS 2026 Assessment',
        'slug': 'arts-2026-assessment'
    },
    {
        'file': 'Qusetion Paper COMMERCE 2026.xlsx',
        'title': 'COMMERCE 2026 Assessment',
        'slug': 'commerce-2026-assessment'
    },
    {
        'file': 'Qusetion Paper SCIENCE 2026.xlsx',
        'title': 'SCIENCE 2026 Assessment',
        'slug': 'science-2026-assessment'
    }
]

def run_step_by_step_verification():
    conn = mysql.connector.connect(**DB_CONFIG)
    cursor = conn.cursor(dictionary=True)
    
    report = {
        'total_assessments_verified': 0,
        'total_parts_verified': 0,
        'total_questions_verified': 0,
        'total_images_verified_on_disk': 0,
        'missing_images_on_disk': [],
        'discrepancies': [],
        'details': []
    }
    
    print("==========================================================================")
    print("STEP-BY-STEP VERIFICATION: ADMIN PANEL / DB vs EXCEL FILES (alpha_tests)")
    print("==========================================================================\n")
    
    for info in excel_files_info:
        fname = info['file']
        title = info['title']
        slug = info['slug']
        fpath = os.path.join(alpha_dir, fname)
        
        print(f"--> STEP 1: Verifying Assessment: '{title}' ({fname})")
        
        # Check assessment in DB
        cursor.execute("SELECT * FROM assessments WHERE title = %s OR slug = %s", (title, slug))
        db_ass = cursor.fetchone()
        
        if not db_ass:
            err = f"CRITICAL: Assessment '{title}' not found in database!"
            report['discrepancies'].append(err)
            print(f"    [FAIL] {err}")
            continue
            
        ass_id = db_ass['id']
        report['total_assessments_verified'] += 1
        print(f"    [OK] DB Assessment ID: {ass_id} | Slug: {db_ass['slug']} | Time: {db_ass['time_limit']} mins")
        
        # Load Excel file
        wb = openpyxl.load_workbook(fpath, data_only=True)
        excel_sheetnames = wb.sheetnames
        
        # Check Parts in DB
        cursor.execute("SELECT * FROM assessment_parts WHERE assessment_id = %s ORDER BY id ASC", (ass_id,))
        db_parts = cursor.fetchall()
        db_part_names = [p['part_name'].strip() for p in db_parts]
        
        print(f"--> STEP 2: Verifying Section/Part Structure ({len(excel_sheetnames)} parts)")
        if len(db_parts) != len(excel_sheetnames):
            err = f"Mismatch in part count for '{title}': Excel has {len(excel_sheetnames)}, DB has {len(db_parts)}"
            report['discrepancies'].append(err)
            print(f"    [FAIL] {err}")
        else:
            print(f"    [OK] Section Count Match: {len(db_parts)} Sections")
            
        ass_detail = {
            'title': title,
            'id': ass_id,
            'parts': []
        }
        
        for sname in excel_sheetnames:
            ws = wb[sname]
            clean_sname = sname.strip()
            
            # Find matching part in DB
            cursor.execute("SELECT * FROM assessment_parts WHERE assessment_id = %s AND TRIM(part_name) = %s", (ass_id, clean_sname))
            db_p = cursor.fetchone()
            
            if not db_p:
                err = f"Part '{clean_sname}' from Excel sheet missing in DB for Assessment {ass_id}!"
                report['discrepancies'].append(err)
                print(f"    [FAIL] {err}")
                continue
                
            part_id = db_p['id']
            report['total_parts_verified'] += 1
            
            # Count questions in Excel sheet
            excel_q_rows = 0
            for r in range(1, ws.max_row + 1):
                val1 = ws.cell(r, 1).value
                val2 = ws.cell(r, 2).value
                for candidate in [val2, val1]:
                    if candidate is not None:
                        try:
                            v_str = str(candidate).strip()
                            if v_str.isdigit() and 1 <= int(v_str) <= 30:
                                excel_q_rows += 1
                                break
                        except ValueError:
                            pass
                            
            # Count questions in DB
            cursor.execute("SELECT * FROM questions WHERE assessment_id = %s AND part_id = %s ORDER BY question_number ASC", (ass_id, part_id))
            db_questions = cursor.fetchall()
            
            print(f"    -> Section [{clean_sname}] (Part ID: {part_id}): Excel Qs = {excel_q_rows}, DB Qs = {len(db_questions)}")
            
            if excel_q_rows != len(db_questions):
                err = f"Question count mismatch in [{clean_sname}]: Excel={excel_q_rows}, DB={len(db_questions)}"
                report['discrepancies'].append(err)
                print(f"       [FAIL] {err}")
            else:
                report['total_questions_verified'] += len(db_questions)
                
            # Verify images on disk for these questions
            part_img_count = 0
            part_opt_img_count = 0
            
            for db_q in db_questions:
                img_fields = [
                    ('image_path', db_q['image_path']),
                    ('image_path_2', db_q['image_path_2']),
                    ('option_a_image', db_q['option_a_image']),
                    ('option_b_image', db_q['option_b_image']),
                    ('option_c_image', db_q['option_c_image']),
                    ('option_d_image', db_q['option_d_image'])
                ]
                
                has_any_img = False
                for field_name, rel_path in img_fields:
                    if rel_path and rel_path.strip() != '':
                        has_any_img = True
                        full_p = os.path.join(fcpath, rel_path.strip())
                        if os.path.exists(full_p):
                            report['total_images_verified_on_disk'] += 1
                        else:
                            err = f"Image file missing on disk: {rel_path} for Q#{db_q['question_number']} in [{clean_sname}]"
                            report['missing_images_on_disk'].append(err)
                            print(f"       [FAIL] {err}")
                            
                if has_any_img:
                    part_img_count += 1
                if any(db_q[f] for f in ['option_a_image', 'option_b_image', 'option_c_image', 'option_d_image']):
                    part_opt_img_count += 1
                    
            ass_detail['parts'].append({
                'part_name': clean_sname,
                'part_id': part_id,
                'questions_count': len(db_questions),
                'questions_with_images': part_img_count,
                'questions_with_option_images': part_opt_img_count
            })
            
        report['details'].append(ass_detail)
        print("")
        
    conn.close()
    
    print("\n==========================================================================")
    print("VERIFICATION SUMMARY:")
    print(f"  Assessments Verified: {report['total_assessments_verified']} / {len(excel_files_info)}")
    print(f"  Sections / Parts Verified: {report['total_parts_verified']} / 36")
    print(f"  Questions Verified: {report['total_questions_verified']} / 720")
    print(f"  Images Verified Intact on Disk: {report['total_images_verified_on_disk']}")
    print(f"  Missing Image Files on Disk: {len(report['missing_images_on_disk'])}")
    print(f"  Total Data Discrepancies: {len(report['discrepancies'])}")
    print("==========================================================================")
    
    with open(r"d:\xampp\htdocs\AlphaMindz\scratch\verification_report.json", "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2, ensure_ascii=False)

if __name__ == '__main__':
    run_step_by_step_verification()

